from __future__ import annotations
from pathlib import Path
import json, html, base64, mimetypes
import cairosvg
from PIL import Image
from pptx import Presentation
from pptx.util import Mm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

WORK=Path(r'C:\Users\ASUS\Documents\ChatGPT\灾害信息处理\scipilot_redraw_20260928')
PANEL_DIR=WORK/'figure7_panels'
meta=json.loads((WORK/'figure7_panels.json').read_text(encoding='utf-8'))
selection=json.loads((Path(r'C:\Users\ASUS\Documents\ChatGPT\灾害信息处理\figures_publication_20260925\Figure9_selection.json')).read_text(encoding='utf-8'))
# Sort by original upper-left positions and then by reading order. The source SVG uses negative y.
meta=sorted(meta,key=lambda r:(r['y'],r['x']))
if len(meta)!=24 or len(selection)!=4: raise RuntimeError('Unexpected panel metadata count')
# D5-2: colourblind-safe error maps (TP light grey, FP dark blue, FN sky blue)
import numpy as _np
CB=WORK/'figure7_panels_colorblind'; CB.mkdir(exist_ok=True)
_REF=_np.array([[255,255,255],[25,178,63],[229,38,38],[25,89,242]],dtype=float)
_NEW=_np.array([[255,255,255],[217,217,217],[8,81,156],[107,174,214]],dtype=float)
_orig=list(meta)
for _i in [5,11,17,23]:
    _src=Path(_orig[_i]['file'])
    _img=_np.asarray(Image.open(_src).convert('RGB'),dtype=float)
    _lab=((_img[:,:,None,:]-_REF[None,None,:,:])**2).sum(-1).argmin(-1)
    _dst=CB/_src.name
    Image.fromarray(_NEW[_lab].astype('uint8')).save(_dst)
    _orig[_i]['file']=str(_dst)
# B3: drop the 'Pooled source' column (its Lombok row carried no information)
_kept=[0,1,2,4,5]
meta=[_orig[r*6+c] for r in range(4) for c in _kept]
cols=['RGB','Ground truth','Source-only','20-shot adapted','Error map']
rows=['Hokkaido\nmedian error','Lombok\nmedian error','Palu\nmedian error','Palu\n90th-percentile failure']
colors={'blue':'#2166AC','red':'#B2182B','green':'#1B7837','orange':'#D88700','ink':'#1F2933','muted':'#5F6B76','light':'#D9DEE3'}
W,H=183.0,123.0
left,right,top,bottom=31.0,2.0,8.0,7.0
gapx,gapy=1.4,1.5
cell_w=(W-left-right-gapx*(len(cols)-1))/len(cols)
cell_h=(H-top-bottom-3.7-gapy*(len(rows)-1))/len(rows)
side=min(cell_w,cell_h)
# center images in cells
img_x=[]
for c in range(len(cols)): img_x.append(left+c*(cell_w+gapx)+(cell_w-side)/2)
img_y=[]
for r in range(len(rows)): img_y.append(top+3.7+r*(cell_h+gapy)+(cell_h-side)/2)

def txt(x,y,s,size=6.0,color=None,weight='normal',anchor='start'):
 return f'<text x="{x:.3f}" y="{y:.3f}" font-family="Arial, Helvetica, sans-serif" font-size="{size:.2f}" fill="{color or colors["ink"]}" font-weight="{weight}" text-anchor="{anchor}">{html.escape(s)}</text>'

svg=[f'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" width="{W}mm" height="{H}mm" viewBox="0 0 {W} {H}">',f'<rect x="0" y="0" width="{W}" height="{H}" fill="white"/>']
# Column headers
for c,name in enumerate(cols): svg.append(txt(img_x[c]+side/2,top+2.0,name,5.4,colors['ink'],'bold','middle'))
# Row labels and images / IoU annotations
for r in range(len(rows)):
 yy=img_y[r]+side/2
 svg.append(txt(1.0,yy-1.2,rows[r].split('\n')[0],5.4,colors['ink'],'bold','start'))
 svg.append(txt(1.0,yy+3.2,rows[r].split('\n')[1],5.1,colors['muted'],'normal','start'))
 for c in range(len(cols)):
  rec=meta[r*len(cols)+c]
  x,y,w,h=img_x[c],img_y[r],side,side
  data = base64.b64encode(Path(rec['file']).read_bytes()).decode('ascii'); href=f"data:image/png;base64,{data}"
  svg.append(f'<rect x="{x:.3f}" y="{y:.3f}" width="{w:.3f}" height="{h:.3f}" fill="white" stroke="{colors["light"]}" stroke-width="0.45"/>')
  svg.append(f'<image x="{x:.3f}" y="{y:.3f}" width="{w:.3f}" height="{h:.3f}" href="{href}" xlink:href="{href}" preserveAspectRatio="xMidYMid meet"/>')
  if c in (2,3):
   iou=[selection[r]['source_iou'],selection[r]['adapted_iou']][c-2]
   svg.append(f'<rect x="{x+w-14.2:.3f}" y="{y+h-4.5:.3f}" width="13.8" height="4.0" rx="0.55" fill="#1F2933" fill-opacity="0.78"/>')
   svg.append(txt(x+w-7.3,y+h-1.75,f'IoU {iou:.3f}',4.55,'#FFFFFF','bold','middle'))
# FP/FN legend
legend_y=H-2.7
svg.append(f'<rect x="{left:.3f}" y="{legend_y-2.0:.3f}" width="3.5" height="3.0" rx="0.35" fill="#D9D9D9"/>')
svg.append(txt(left+4.4,legend_y,'true positive',4.9,colors['muted']))
svg.append(f'<rect x="{left+34:.3f}" y="{legend_y-2.0:.3f}" width="3.5" height="3.0" rx="0.35" fill="#08519C"/>')
svg.append(txt(left+38.4,legend_y,'false positive',4.9,colors['muted']))
svg.append(f'<rect x="{left+73:.3f}" y="{legend_y-2.0:.3f}" width="3.5" height="3.0" rx="0.35" fill="#6BAED6"/>')
svg.append(txt(left+77.4,legend_y,'false negative',4.9,colors['muted']))
svg.append(txt(W-right,legend_y,'Data panels are separate raster observations; labels, frames and legend are vector.',4.5,colors['muted'],'normal','end'))
svg.append('</svg>')
svg_text='\n'.join(svg)
svg_path=WORK/'Figure7_final.svg'; svg_path.write_text(svg_text,encoding='utf-8')
cairosvg.svg2pdf(url=str(svg_path),write_to=str(WORK/'Figure7_final.pdf'),output_width=W*4,output_height=H*4)
cairosvg.svg2png(url=str(svg_path),write_to=str(WORK/'Figure7_final.png'),output_width=W*14,output_height=H*14)

# Native PowerPoint version: individual image panels remain separate pictures, all labels and frames are native vector shapes.
prs=Presentation(); prs.slide_width=Mm(W); prs.slide_height=Mm(H); slide=prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb=RGBColor.from_string('FFFFFF')
def rgb(x): return RGBColor.from_string(x.lstrip('#'))
def add_text(x,y,w,h,s,size,color,bold=False,align=PP_ALIGN.LEFT):
 box=slide.shapes.add_textbox(Mm(x),Mm(y),Mm(w),Mm(h)); tf=box.text_frame; tf.clear(); tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0; tf.vertical_anchor=MSO_ANCHOR.TOP; tf.word_wrap=True
 for i,line in enumerate(s.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line; p.alignment=align; p.space_before=Pt(0); p.space_after=Pt(0)
  for run in p.runs: run.font.name='Arial'; run.font.size=Pt(size); run.font.bold=bold; run.font.color.rgb=rgb(color)
 return box
for c,name in enumerate(cols): add_text(img_x[c],0.9,cell_w,2.6,name,5.4,colors['ink'],True,PP_ALIGN.CENTER)
for r in range(len(rows)):
 add_text(1.0,img_y[r],28.0,3.0,rows[r].split('\n')[0],5.4,colors['ink'],True)
 add_text(1.0,img_y[r]+3.3,28.0,3.0,rows[r].split('\n')[1],5.1,colors['muted'])
 for c in range(len(cols)):
  rec=meta[r*len(cols)+c]; path=Path(rec['file']); slide.shapes.add_picture(str(path),Mm(img_x[c]),Mm(img_y[r]),Mm(side),Mm(side))
  frame=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Mm(img_x[c]),Mm(img_y[r]),Mm(side),Mm(side)); frame.fill.background(); frame.line.color.rgb=rgb(colors['light']); frame.line.width=Pt(0.35)
  if c in (2,3):
   iou=[selection[r]['source_iou'],selection[r]['adapted_iou']][c-2]
   bg=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Mm(img_x[c]+side-14.2),Mm(img_y[r]+side-4.5),Mm(13.8),Mm(4.0)); bg.fill.solid(); bg.fill.fore_color.rgb=rgb(colors['ink']); bg.line.fill.background()
   add_text(img_x[c]+side-14.2,img_y[r]+side-3.8,13.8,2.3,f'IoU {iou:.3f}',4.55,'#FFFFFF',True,PP_ALIGN.CENTER)
legend_y=H-3.0
for idx,label,color in [(0,'true positive','#D9D9D9'),(1,'false positive','#08519C'),(2,'false negative','#6BAED6')]:
 x=left+idx*39.0
 shp=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Mm(x),Mm(legend_y),Mm(3.5),Mm(3.0)); shp.fill.solid(); shp.fill.fore_color.rgb=rgb(color); shp.line.fill.background()
 add_text(x+4.4,legend_y+0.1,30,3,label,4.9,colors['muted'])
add_text(left+120,legend_y+0.1,60,3,'Data panels are separate observations.',4.5,colors['muted'],align=PP_ALIGN.RIGHT)
prs.save(str(WORK/'Figure7_final.pptx'))
print(svg_path)
