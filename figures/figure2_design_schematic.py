# -*- coding: utf-8 -*-
"""Figure 2 v3 - MDPI Remote Sensing style experimental-architecture schematic (fully vector)."""
import os
from xml.sax.saxutils import escape

W, H = 510.24, 336.0
INK="#1F2A33"; FRAME="#8A94A0"; FRAME_FILL="#FBFCFD"; SUB="#6B7785"
REF="#1F6FB2"; GATE="#C0831B"; PRIM="#B3352F"; DIAG="#2E7D5B"
T_REF="#E9F1FA"; T_GATE="#FBF3E3"; T_PRIM="#FAECEA"; T_DIAG="#EAF4EF"; T_SUB="#F1F4F7"

OPS=[]
def box(x,y,w,h,stroke,fill,width=0.9,dash=None,r=2.0): OPS.append(("box",x,y,w,h,stroke,fill,width,dash,r))
def circle(cx,cy,rad,stroke,fill,width=0.8): OPS.append(("circle",cx,cy,rad,stroke,fill,width))
def line(x1,y1,x2,y2,c=SUB,width=0.9): OPS.append(("line",x1,y1,x2,y2,c,width))
def text(x,y,s,size=5.8,bold=False,c=INK,anchor="start"): OPS.append(("text",x,y,s,size,bold,c,anchor))
def arrowR(x,y,w,c=SUB): OPS.append(("arrowR",x,y,w,c))
def arrowD(x,y,h,c=SUB): OPS.append(("arrowD",x,y,h,c))

def icon(kind,x,y,s,c):
    if kind=="stack":
        for dx,dy in [(0.00,0.34),(0.17,0.17),(0.34,0.00)]:
            box(x+dx*s, y+dy*s, 0.58*s, 0.5*s, c, "#FFFFFF", 0.7, r=1.0)
    elif kind=="grid":
        for dx,dy in [(0.02,0.02),(0.52,0.02),(0.02,0.52),(0.52,0.52)]:
            box(x+dx*s, y+dy*s, 0.42*s, 0.42*s, c, "#FFFFFF", 0.7, r=1.0)
    elif kind=="ring":
        circle(x+0.5*s, y+0.5*s, 0.42*s, c, "#FFFFFF", 0.9)
        circle(x+0.5*s, y+0.5*s, 0.16*s, c, c, 0.6)
    elif kind=="bars":
        for dx,hh in [(0.10,0.45),(0.40,0.70),(0.70,1.00)]:
            box(x+dx*s, y+(1.0-hh)*s, 0.20*s, hh*s, c, c, 0.5, r=0.6)
    elif kind=="lock":
        box(x+0.14*s, y+0.42*s, 0.72*s, 0.56*s, c, "#FFFFFF", 0.8, r=1.0)
        line(x+0.30*s, y+0.42*s, x+0.30*s, y+0.22*s, c, 0.9)
        line(x+0.70*s, y+0.42*s, x+0.70*s, y+0.22*s, c, 0.9)
        line(x+0.30*s, y+0.22*s, x+0.70*s, y+0.22*s, c, 0.9)
    elif kind=="split":
        box(x+0.06*s, y+0.34*s, 0.40*s, 0.50*s, c, c, 0.5, r=1.0)
        box(x+0.54*s, y+0.34*s, 0.40*s, 0.50*s, c, "#FFFFFF", 0.8, r=1.0)
        line(x+0.06*s, y+0.16*s, x+0.94*s, y+0.16*s, c, 0.9)
    elif kind=="nodes":
        circle(x+0.16*s, y+0.50*s, 0.15*s, c, "#FFFFFF", 0.8)
        circle(x+0.84*s, y+0.22*s, 0.15*s, c, "#FFFFFF", 0.8)
        circle(x+0.84*s, y+0.78*s, 0.15*s, c, "#FFFFFF", 0.8)
        line(x+0.30*s, y+0.46*s, x+0.70*s, y+0.28*s, c, 0.6)
        line(x+0.30*s, y+0.54*s, x+0.70*s, y+0.72*s, c, 0.6)

# ---------- Frame A ----------
box(6,6,498.24,72, FRAME, FRAME_FILL, 1.1, dash=True, r=4)
circle(20,17.5,6.2, REF, REF, 0.7); text(20,19.6,"1",4.6,True,"#FFFFFF","middle")
text(31,20,"Data and study units",8.0,True,INK)
bw=155.41; bxs=[14,177.41,340.82]
data=[("stack",REF,T_REF,"Source domain (training)",
       ["CAS Moxi + Bijie; 1,795 tiles;","50 epochs; checkpoint selection","uses source validation only."]),
      ("grid",GATE,T_GATE,"Added-source pool (5 slots, equal mass)",
       ["CAS, Jiuzhaigou, Longxi River,","Moxitaidi, Wenchuan; 11,382 tiles;","equal slot mass (7,168 exposures)."]),
      ("ring",SUB,T_SUB,"Held-out target units (n = 6)",
       ["Primary: Hokkaido, Lombok, Palu.","Exploratory LORO: Wenchuan,","Jiuzhaigou, Longxi River.","Moxitaidi: source-adjacent."])]
for x,(ic,ac,tn,title,lines) in zip(bxs,data):
    box(x,24,bw,48, ac, "#FFFFFF", 0.9, r=2.0)
    box(x,24,3.0,48, ac, ac, 0.5, r=0.0)
    icon(ic, x+8, 28, 14, ac)
    text(x+27, 37, title, 6.3, True, INK)
    yy=46.5
    for ln in lines:
        text(x+8, yy, ln, 5.6, False, SUB); yy+=7.4
arrowD(W/2, 78, 8)
# ---------- Frame B ----------
box(6,86,498.24,142, FRAME, FRAME_FILL, 1.1, dash=True, r=4)
circle(20,97.5,6.2, GATE, GATE, 0.7); text(20,99.6,"2",4.6,True,"#FFFFFF","middle")
text(31,100,"Experimental design (four controls)",8.0,True,INK)
cw=115.0; cxs=[14,139,264,389]
cards=[("bars",REF,T_REF,"REFERENCE","1. Reference benchmark",
        ["Source-only training;","4 architectures x 3 seeds","(12 source-only runs).","Gives the initial checkpoints","and the zero-shot baseline.","Report: Table 4; Figure 3."]),
       ("lock",GATE,T_GATE,"CONFIGURATION GATE","2. Configuration lock (X2)",
        ["current / none / strong;","ResUNet + Bottleneck-LiteASK;","3 seeds; 18 paired runs.","Locked by source-validation","IoU only, before any target","result. Selected: none.","Report: Table 5."]),
       ("split",PRIM,T_PRIM,"PRIMARY CONTRAST","3. Pooled-source contrast",
        ["Exposure-matched: single stream,","five-slot CAS replay placebo and","five-slot equal-mass pooled source.","1,120 steps; 35,840 exposures.","Epoch-matched variant:","exploratory only.","Report: Tables 6-8; Figure 4a,b."]),
       ("nodes",DIAG,T_DIAG,"DIAGNOSTIC","4. Diagnostics",
        ["Resolution (X1): 128, 256 and","128-weight-at-256 arms.","Few-shot adaptation: 5/10/20","shots; physical-buffer and","support-draw checks.","Report: Table 9; Figure 4c,d;","Figures 5-7."])]
for x,(ic,ac,tn,tag,title,lines) in zip(cxs,cards):
    box(x,106,cw,116, ac, "#FFFFFF", 0.9, r=2.0)
    box(x,106,3.0,116, ac, ac, 0.5, r=0.0)
    circle(x+12,120,6.4, ac, ac, 0.6); text(x+12,122.2,title[0],4.8,True,"#FFFFFF","middle")
    OPS.append(("pill",x+23,112.5,cw-31,10.5,ac,tn,tag))
    text(x+8, 134, title, 6.3, True, INK)
    icon(ic, x+cw-22, 127, 14, ac)
    yy=145
    for ln in lines:
        text(x+8, yy, ln, 5.5, False, SUB); yy+=7.6
for i in range(3):
    arrowR(cxs[i]+cw+1.4, 164, 10-2.8)
arrowD(W/2, 228, 8)
# ---------- Frame C ----------
box(6,236,498.24,58, FRAME, FRAME_FILL, 1.1, dash=True, r=4)
circle(20,247.5,6.2, DIAG, DIAG, 0.7); text(20,249.6,"3",4.6,True,"#FFFFFF","middle")
text(31,250,"Evaluation and inference",8.0,True,INK)
ev=[("Endpoints",["IoU (primary), balanced IoU,","MCC, precision/recall; BF1 at","2 and 4 px, HD95 (few-shot)."]),
    ("Reporting",["Region-level reporting is","mandatory; all six held-out","units reported separately."]),
    ("Uncertainty",["Seed-stratified Student t","across three seeds; connected-","component bootstrap in seed."]),
    ("Operating point",["Thresholds 0.3-0.7; retention =","cross-event macro IoU /","source-validation IoU."])]
for x,(title,lines) in zip(cxs,ev):
    box(x,254,cw,34, SUB, "#FFFFFF", 0.7, r=1.5)
    text(x+6, 263, title, 5.9, True, INK)
    yy=272
    for ln in lines:
        text(x+6, yy, ln, 5.3, False, SUB); yy+=7.0
# ---------- legend + footnotes ----------
leg=[(REF,"Reference benchmark"),(GATE,"Configuration gate"),(PRIM,"Primary contrast"),(DIAG,"Diagnostic"),(SUB,"Exploratory (unmatched exposure)")]
step=(W-2*7)/5.0
for i,(c,lab) in enumerate(leg):
    x=7+i*step
    box(x,301,10,8,c,c,0.6,r=1.5); text(x+13.5,308,lab,5.6,False,INK)
line(7,313,W-7,313,FRAME,0.7)
text(7,321,"Solid arrows: primary experimental flow. The epoch-matched comparison is exploratory and is not used for the primary claim.",5.5,False,SUB)
text(7,329,"Experiment names, region roles and evidence boundaries follow the fixed pipeline in Sections 4.2-4.6.",5.5,False,SUB)

# ---------- renderers ----------
def render_svg(path):
    p=['<?xml version="1.0" encoding="utf-8"?>',
       '<svg xmlns="http://www.w3.org/2000/svg" width="%.2fpt" height="%.2fpt" viewBox="0 0 %.2f %.2f">'%(W,H,W,H),
       '<rect x="0" y="0" width="%.2f" height="%.2f" fill="#FFFFFF"/>'%(W,H)]
    for op in OPS:
        k=op[0]
        if k=="box":
            _,x,y,w,h,sc,fc,wd,dash,r=op
            d=' stroke-dasharray="4,2.5"' if dash else ''
            p.append('<rect x="%.2f" y="%.2f" width="%.2f" height="%.2f" rx="%.2f" fill="%s" stroke="%s" stroke-width="%.2f"%s/>'%(x,y,w,h,r,fc,sc,wd,d))
        elif k=="circle":
            _,cx,cy,rad,sc,fc,wd=op
            p.append('<circle cx="%.2f" cy="%.2f" r="%.2f" fill="%s" stroke="%s" stroke-width="%.2f"/>'%(cx,cy,rad,fc,sc,wd))
        elif k=="line":
            _,x1,y1,x2,y2,c,wd=op
            p.append('<line x1="%.2f" y1="%.2f" x2="%.2f" y2="%.2f" stroke="%s" stroke-width="%.2f"/>'%(x1,y1,x2,y2,c,wd))
        elif k=="arrowR":
            _,x,y,w,c=op; x2=x+w
            p.append('<line x1="%.2f" y1="%.2f" x2="%.2f" y2="%.2f" stroke="%s" stroke-width="1.1"/>'%(x,y,x2-3,y,c))
            p.append('<path d="M%.2f %.2f L%.2f %.2f L%.2f %.2f Z" fill="%s"/>'%(x2-3,y-2.6,x2,y,x2-3,y+2.6,c))
        elif k=="arrowD":
            _,x,y,h,c=op; y2=y+h
            p.append('<line x1="%.2f" y1="%.2f" x2="%.2f" y2="%.2f" stroke="%s" stroke-width="1.1"/>'%(x,y,x,y2-3,c))
            p.append('<path d="M%.2f %.2f L%.2f %.2f L%.2f %.2f Z" fill="%s"/>'%(x-2.6,y2-3,x,y2,x+2.6,y2-3,c))
        elif k=="pill":
            _,x,y,w,h,ac,tn,tag=op
            p.append('<rect x="%.2f" y="%.2f" width="%.2f" height="%.2f" rx="%.2f" fill="%s"/>'%(x,y,w,h,h/2,tn))
            p.append('<text x="%.2f" y="%.2f" font-family="Arial, Helvetica, sans-serif" font-size="4.8" font-weight="700" fill="%s">%s</text>'%(x+5,y+h-3.2,ac,escape(tag)))
        elif k=="text":
            _,x,y,s,size,b,cc,anchor=op
            p.append('<text x="%.2f" y="%.2f" font-family="Arial, Helvetica, sans-serif" font-size="%.2f" font-weight="%s" fill="%s" text-anchor="%s">%s</text>'%(x,y,size,"700" if b else "400",cc,anchor,escape(s)))
    p.append('</svg>')
    open(path,"w",encoding="utf-8").write("\n".join(p))

def render_pdf(path):
    from reportlab.pdfgen import canvas as rlcanvas
    c=rlcanvas.Canvas(path,pagesize=(W,H))
    def Y(y): return H-y
    for op in OPS:
        k=op[0]
        if k=="box":
            _,x,y,w,h,sc,fc,wd,dash,r=op
            c.setFillColor(fc); c.setStrokeColor(sc); c.setLineWidth(wd)
            c.setDash([4,2.5] if dash else [])
            c.roundRect(x,Y(y+h),w,h,r,stroke=1,fill=1); c.setDash([])
        elif k=="circle":
            _,cx,cy,rad,sc,fc,wd=op
            c.setFillColor(fc); c.setStrokeColor(sc); c.setLineWidth(wd)
            c.circle(cx,Y(cy),rad,stroke=1,fill=1)
        elif k=="line":
            _,x1,y1,x2,y2,cc,wd=op
            c.setStrokeColor(cc); c.setLineWidth(wd); c.line(x1,Y(y1),x2,Y(y2))
        elif k=="arrowR":
            _,x,y,w,cc=op; x2=x+w
            c.setStrokeColor(cc); c.setFillColor(cc); c.setLineWidth(1.1)
            c.line(x,Y(y),x2-3,Y(y))
            pa=c.beginPath(); pa.moveTo(x2-3,Y(y+2.6)); pa.lineTo(x2,Y(y)); pa.lineTo(x2-3,Y(y-2.6)); pa.close(); c.drawPath(pa,stroke=0,fill=1)
        elif k=="arrowD":
            _,x,y,h,cc=op; y2=y+h
            c.setStrokeColor(cc); c.setFillColor(cc); c.setLineWidth(1.1)
            c.line(x,Y(y),x,Y(y2-3))
            pa=c.beginPath(); pa.moveTo(x-2.6,Y(y2-3)); pa.lineTo(x,Y(y2)); pa.lineTo(x+2.6,Y(y2-3)); pa.close(); c.drawPath(pa,stroke=0,fill=1)
        elif k=="pill":
            _,x,y,w,h,ac,tn,tag=op
            c.setFillColor(tn); c.setStrokeColor(tn); c.setLineWidth(0.4)
            c.roundRect(x,Y(y+h),w,h,h/2,stroke=1,fill=1)
            c.setFillColor(ac); c.setFont("Helvetica-Bold",4.8); c.drawString(x+5,Y(y+h-3.2),tag)
        elif k=="text":
            _,x,y,s,size,b,cc,anchor=op
            c.setFillColor(cc); c.setFont("Helvetica-Bold" if b else "Helvetica",size)
            if anchor=="end": c.drawRightString(x,Y(y),s)
            elif anchor=="middle": c.drawCentredString(x,Y(y),s)
            else: c.drawString(x,Y(y),s)
    c.showPage(); c.save()

def render_pptx(path):
    from pptx import Presentation
    from pptx.util import Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    prs=Presentation(); prs.slide_width=Emu(int(W*12700)); prs.slide_height=Emu(int(H*12700))
    sl=prs.slides.add_slide(prs.slide_layouts[6])
    def C(h): return RGBColor.from_string(h.lstrip("#").upper())
    def E(v): return Emu(int(round(v*12700)))
    def T(x,y,s,size,b,c,anchor):
        w=max(20.0,4.7*len(s)+6)
        off = w if anchor=="end" else (w/2 if anchor=="middle" else 0)
        tb=sl.shapes.add_textbox(E(x-off), E(y-size*0.85), E(w), E(size*1.55))
        tf=tb.text_frame; tf.word_wrap=False
        tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
        p=tf.paragraphs[0]; p.alignment=PP_ALIGN.RIGHT if anchor=="end" else (PP_ALIGN.CENTER if anchor=="middle" else PP_ALIGN.LEFT)
        r=p.add_run(); r.text=s; r.font.size=Pt(size); r.font.bold=b; r.font.name="Arial"; r.font.color.rgb=C(c)
    for op in OPS:
        k=op[0]
        if k=="box":
            _,x,y,w,h,sc,fc,wd,dash,r=op
            sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,E(x),E(y),E(w),E(h))
            try: sh.adjustments[0]=max(0.0,min(0.5,r/min(w,h)))
            except Exception: pass
            sh.fill.solid(); sh.fill.fore_color.rgb=C(fc); sh.line.color.rgb=C(sc); sh.line.width=Pt(wd); sh.shadow.inherit=False
            if dash:
                from pptx.oxml.ns import qn
                ln=sh.line._get_or_add_ln(); ln.append(ln.makeelement(qn('a:prstDash'),{'val':'dash'}))
            sh.text_frame.text=""
        elif k=="circle":
            _,cx,cy,rad,sc,fc,wd=op
            sh=sl.shapes.add_shape(MSO_SHAPE.OVAL,E(cx-rad),E(cy-rad),E(2*rad),E(2*rad))
            sh.fill.solid(); sh.fill.fore_color.rgb=C(fc); sh.line.color.rgb=C(sc); sh.line.width=Pt(wd); sh.shadow.inherit=False
        elif k=="line":
            _,x1,y1,x2,y2,cc,wd=op
            cn=sl.shapes.add_connector(1,E(x1),E(y1),E(x2),E(y2)); cn.line.color.rgb=C(cc); cn.line.width=Pt(wd)
        elif k=="arrowR":
            _,x,y,w,cc=op
            sh=sl.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,E(x),E(y-3),E(w),E(6))
            sh.fill.solid(); sh.fill.fore_color.rgb=C(cc); sh.line.fill.background(); sh.shadow.inherit=False
        elif k=="arrowD":
            _,x,y,h,cc=op
            sh=sl.shapes.add_shape(MSO_SHAPE.DOWN_ARROW,E(x-3),E(y),E(6),E(h))
            sh.fill.solid(); sh.fill.fore_color.rgb=C(cc); sh.line.fill.background(); sh.shadow.inherit=False
        elif k=="pill":
            _,x,y,w,h,ac,tn,tag=op
            sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,E(x),E(y),E(w),E(h))
            try: sh.adjustments[0]=0.5
            except Exception: pass
            sh.fill.solid(); sh.fill.fore_color.rgb=C(tn); sh.line.color.rgb=C(tn); sh.line.width=Pt(0.4); sh.shadow.inherit=False
            tf=sh.text_frame; tf.word_wrap=False
            tf.margin_left=Emu(0); tf.margin_right=Emu(0); tf.margin_top=Emu(0); tf.margin_bottom=Emu(0)
            p0=tf.paragraphs[0]; p0.alignment=PP_ALIGN.LEFT
            r0=p0.add_run(); r0.text=tag; r0.font.size=Pt(4.8); r0.font.bold=True; r0.font.name="Arial"; r0.font.color.rgb=C(ac)
        elif k=="text":
            T(*op[1:])
    prs.save(path)

OUT=r"C:\Users\ASUS\Documents\ChatGPT\灾害信息处理\_revision_work\mdpi_rs_pass\figure2"
os.makedirs(OUT,exist_ok=True)
render_svg(os.path.join(OUT,"Figure 2.svg"))
render_pdf(os.path.join(OUT,"Figure 2.pdf"))
render_pptx(os.path.join(OUT,"Figure 2.pptx"))
print("built v3 (vector only)")
